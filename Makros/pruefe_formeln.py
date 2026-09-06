#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Formelkonsistenz in den fertigen Arbeitsmappen (Pruefebene 6).

WAS DIESE PRUEFUNG FINDEN SOLL
    Der Ursprungsfehler des ganzen Projekts: Spalte B war eine einzelne
    Ueberlauf-Matrixformel, C/D/L/P/U waren je Zelle als Matrixformel
    markiert - und Excel erweitert die beim Zeileneinfuegen NICHT. Nach
    jedem Einfuegen fehlten in der neuen Zeile die Formeln, oder sie
    zeigten auf die falsche Zeile. Sichtbar wurde das erst Wochen spaeter.

WARUM NICHT DIE VBA-FORMELN NACHBAUEN
    Naheliegend waere, die Frm*-Funktionen aus modWochenplan in Python
    nachzubauen und zeichengenau zu vergleichen. Genau das habe ich beim
    Entwickeln von Hand gemacht (309 bzw. 340 Formeln, 0 Abweichungen).
    Als DAUERTEST taugt es nicht: es waere eine zweite Umsetzung
    derselben Idee, die mit jeder VBA-Aenderung neu nachgezogen werden
    muesste - und wer beides aus derselben Vorstellung schreibt, bekommt
    zwei Mal denselben Fehler.

    Diese Pruefung stellt deshalb eine andere Frage, die ohne Kenntnis
    der VBA-Logik auskommt:
        "Sehen alle Planzeilen gleich aus?"
    Eine Zeile, die aus der Reihe faellt, ist immer ein Fehler - egal
    welche Formel richtig waere. Das ist unabhaengig vom Code und genau
    deshalb aussagekraeftig.

AUFRUF
    python3 pruefe_formeln.py ../*.xlsm
    python3 pruefe_formeln.py --selbsttest <eine gefuellte Datei>
"""
import sys, re, argparse, warnings, os, zipfile, shutil, tempfile
warnings.filterwarnings("ignore")
import openpyxl

WP, LB = "Wochenplan", "Lernbereiche"

#  Spalten des Wochenplans, in denen eine Formel stehen MUSS und in denen
#  alle Inhaltszeilen dieselbe Formel tragen (bis auf die Zeilennummern).
WP_FORMELSPALTEN = ["A", "C", "D", "L", "P", "Q", "R", "V"]
#  B faellt raus: nach dem Fixieren stehen dort feste Werte.
#  S faellt raus: es verweist auf die vorherige INHALTSzeile, die bei
#  Ferienzeilen dazwischen unterschiedlich weit entfernt liegt.
#  U faellt raus: Matrixformel mit Metadatenreferenz, openpyxl liest sie
#  nicht zuverlaessig zurueck.
LB_FORMELSPALTEN = ["F", "G", "H", "I"]


def txt(ws, r, c):
    v = ws.cell(r, c).value
    return "" if v is None else str(v).strip()


def kopfzeile(ws, spalte, text):
    for r in range(1, 60):
        if txt(ws, r, spalte).lower().startswith(text.lower()):
            return r
    return 0


def ist_ferienzeile(ws, r, erste):
    if txt(ws, r, 14).upper() in ("F", "FERIEN"):
        return True
    for rng in ws.merged_cells.ranges:
        if (rng.min_row == r and rng.min_col == 2
                and (rng.max_col - rng.min_col + 1) > 4 and r >= erste):
            return True
    return False


def normiere(formel, zeile, erste):
    """Zeilenbezuege durch Platzhalter ersetzen, damit sich zwei Zeilen
    vergleichen lassen.

    ZWEI FALLEN, beide beim ersten Lauf gegen die echten Dateien
    aufgeschlagen - die Pruefung meldete acht Fehler in jeder bekannt
    guten Datei:

    1. Die Ersetzungspaare lagen in einem set, und ein set hat keine
       feste Reihenfolge. In der ersten Planzeile sind zeile und erste
       dieselbe Zahl; je nach Laune der Iteration wurde daraus @Z oder
       @E, und die Musterzeile passte zu keiner anderen.
    2. Eine blanke Zahlenersetzung trifft zu viel: in Zeile 15 wurde aus
       "Einstellungen!$B$15" ein "$B$@Z", in Zeile 10 aus "CHAR(10)" ein
       "CHAR(@Z)", in Zeile 9 aus ">=9" ein ">=@Z".

    Deshalb jetzt zweistufig:
    - Fremdblattbezuege komplett ausblenden. Die sind in jeder Zeile
      absolut und identisch, tragen zum Vergleich also nichts bei.
    - Im Rest nur Zahlen ersetzen, die DIREKT auf einen Spaltenbuchstaben
      folgen - also echte Zellbezuege wie $E4, P4:S4, $V4. "CHAR(10)"
      und ">=9" bleiben unangetastet.
    """
    formel = re.sub(r"[A-Za-z_][A-Za-z0-9_]*!\$?[A-Z]{1,3}\$?\d+"
                    r"(?::\$?[A-Z]{1,3}\$?\d+)?", "@FREMD", formel)

    paare = [("%d" % zeile, "@Z")]
    if erste != zeile:
        paare.append(("%d" % erste, "@E"))
    for zahl, marke in paare:
        formel = re.sub(r"(?<=[A-Z])(\$?)" + zahl + r"(?!\d)",
                        lambda m: m.group(1) + marke, formel)
    return formel


def pruefe(pfad, wb=None):
    fund = []

    def melde(regel, text):
        fund.append((regel, text))

    if wb is None:
        wb = openpyxl.load_workbook(pfad, data_only=False)
    if WP not in wb.sheetnames:
        melde("blaetter", "Blatt %r fehlt" % WP)
        return fund

    ws = wb[WP]
    kopf = kopfzeile(ws, 2, "Referenz Code")
    if kopf == 0:
        melde("blaetter", "Ueberschriftenzeile des Wochenplans nicht gefunden")
        return fund
    erste = kopf + 1

    # letzte Zeile mit Inhalt
    letzte = erste
    for r in range(erste, ws.max_row + 1):
        if any(txt(ws, r, c) for c in (5, 6, 7, 8, 9, 11, 13, 14)):
            letzte = r
    inhalt = [r for r in range(erste, letzte + 1)
              if not ist_ferienzeile(ws, r, erste)]

    if len(inhalt) < 2:
        melde("hinweis", "weniger als zwei Inhaltszeilen - nichts zu vergleichen")
        return fund

    # --- 1. Wochenplan: gleiche Formel in allen Inhaltszeilen --------
    for sp in WP_FORMELSPALTEN:
        c = openpyxl.utils.column_index_from_string(sp)
        muster, musterzeile = None, None
        ohne = []
        for r in inhalt:
            v = ws.cell(r, c).value
            if not (isinstance(v, str) and v.startswith("=")):
                ohne.append(r)
                continue
            n = normiere(v, r, erste)
            if muster is None:
                muster, musterzeile = n, r
            elif n != muster:
                melde("formel-abweichung",
                      "Spalte %s: Zeile %d weicht von Zeile %d ab\n"
                      "          Zeile %d: %s\n"
                      "          Zeile %d: %s"
                      % (sp, r, musterzeile, musterzeile,
                         ws.cell(musterzeile, c).value, r, v))
                break
        if ohne:
            melde("formel-fehlt",
                  "Spalte %s: keine Formel in Zeile %s"
                  % (sp, ohne[:8]))

    # --- 2. Ferienzeilen tragen KEINE Formeln in C ------------------
    #  Sonst rechnet die verbundene Zelle mit und liefert Muell.
    for r in range(erste, letzte + 1):
        if ist_ferienzeile(ws, r, erste):
            v = ws.cell(r, 3).value
            if isinstance(v, str) and v.startswith("="):
                melde("ferienzeile",
                      "Ferienzeile %d hat eine Formel in Spalte C" % r)

    # --- 3. Lernbereiche: dasselbe Spiel ----------------------------
    if LB in wb.sheetnames:
        wl = wb[LB]
        kopfL = kopfzeile(wl, 1, "Referenzcode")
        if kopfL:
            ersteL = kopfL + 1
            zSum = 0
            for r in range(ersteL, ersteL + 80):
                if txt(wl, r, 4).lower() == "summe":
                    zSum = r
                    break
            letzteL = ersteL - 1
            for r in range(ersteL, zSum or ersteL):
                if txt(wl, r, 1) or txt(wl, r, 4):
                    letzteL = r
            if letzteL >= ersteL:
                for sp in LB_FORMELSPALTEN:
                    c = openpyxl.utils.column_index_from_string(sp)
                    muster, musterzeile = None, None
                    for r in range(ersteL, letzteL + 1):
                        v = wl.cell(r, c).value
                        if not (isinstance(v, str) and v.startswith("=")):
                            melde("formel-fehlt",
                                  "%s Spalte %s: keine Formel in Zeile %d"
                                  % (LB, sp, r))
                            continue
                        n = normiere(v, r, ersteL)
                        if muster is None:
                            muster, musterzeile = n, r
                        elif n != muster:
                            melde("formel-abweichung",
                                  "%s Spalte %s: Zeile %d weicht von Zeile %d ab"
                                  % (LB, sp, r, musterzeile))
                            break

    # --- 4. Fremdblattbezuege zeigen auf den richtigen Bereich -------
    #  Der FILTER in Spalte U greift auf einen festen Bereich in
    #  "Lernbereiche" zu. Zeigt er noch auf den Stand vor einer
    #  Verschiebung, bleibt die Spalte still LEER - ohne Fehlermeldung,
    #  ohne #REF!. Genau dieser Fehler steckte einmal drin
    #  ($A$2:$A$50 statt $A$4:$A$52).
    #
    #  Gelesen wird aus der Roh-XML, nicht ueber openpyxl: Spalte U ist
    #  eine Matrixformel mit Metadatenreferenz, die openpyxl nicht
    #  zurueckliefert. Ein Check darauf waere stumm geblieben - und eine
    #  stumme Pruefung ist schlimmer als keine.
    if pfad and LB in wb.sheetnames:
        kopfL = kopfzeile(wb[LB], 1, "Referenzcode")
        if kopfL:
            gefunden = set()
            try:
                with zipfile.ZipFile(pfad) as z:
                    for name in z.namelist():
                        if not name.startswith("xl/worksheets/sheet"):
                            continue
                        roh = z.read(name).decode("utf-8", "replace")
                        for m in re.finditer(r"Lernbereiche!\$A\$(\d+)", roh):
                            gefunden.add(int(m.group(1)))
            except Exception as e:
                melde("bezug", "XML nicht lesbar: %r" % e)
            if gefunden and min(gefunden) != kopfL + 1:
                melde("bezug",
                      "Bezug auf Lernbereiche!$A$%d, erste Datenzeile der "
                      "%s ist aber %d" % (min(gefunden), LB, kopfL + 1))

    return fund


# ------------------------------------------------------- Selbstpruefung
def selbsttest(pfad):
    basis = [f for f in pruefe(pfad) if f[0] != "hinweis"]
    print("Ausgangslage: %d Befund(e)" % len(basis))
    for r, s in basis:
        print("   [%s] %s" % (r, s.split("\n")[0]))
    if basis:
        print("\nACHTUNG: Die Ausgangsdatei ist nicht sauber. Eine Regel, die "
              "hier schon\nrot ist, kann durch eine Mutation nichts mehr "
              "beweisen.")
    vorhanden = {r for r, _ in basis}

    def m_formel_kaputt(wb):
        """Eine einzelne Formel verbiegen - der Matrixformel-Fehler."""
        ws = wb[WP]
        kopf = kopfzeile(ws, 2, "Referenz Code")
        for r in range(kopf + 1, ws.max_row + 1):
            v = ws.cell(r, 1).value
            if isinstance(v, str) and v.startswith("="):
                ws.cell(r, 1).value = v.replace('="', '=IF(TRUE,"', 1) + ")"
                return

    def m_formel_weg(wb):
        """Formel durch einen festen Wert ersetzt."""
        ws = wb[WP]
        kopf = kopfzeile(ws, 2, "Referenz Code")
        for r in range(kopf + 2, ws.max_row + 1):
            v = ws.cell(r, 3).value
            if isinstance(v, str) and v.startswith("="):
                ws.cell(r, 3).value = "01.01."
                return

    def m_lb_formel(wb):
        wl = wb[LB]
        k = kopfzeile(wl, 1, "Referenzcode")
        for r in range(k + 2, k + 40):
            v = wl.cell(r, 6).value
            if isinstance(v, str) and v.startswith("="):
                wl.cell(r, 6).value = "=SUM(1)"
                return

    def m_bezug_datei(quelle):
        """Verstellt den FILTER-Bezug auf XML-Ebene und liefert eine
        neue Datei - openpyxl kann diese Formel nicht schreiben."""
        ziel = tempfile.mktemp(suffix=".xlsm")
        with zipfile.ZipFile(quelle) as zin, \
             zipfile.ZipFile(ziel, "w", zipfile.ZIP_DEFLATED) as zout:
            for it in zin.infolist():
                d = zin.read(it.filename)
                if it.filename.startswith("xl/worksheets/sheet"):
                    d = re.sub(rb"Lernbereiche!\$A\$\d+",
                               b"Lernbereiche!$A$99", d)
                zout.writestr(it, d)
        return ziel

    faelle = [
        ("Formel in einer Zeile verbogen", "formel-abweichung", m_formel_kaputt),
        ("Formel durch festen Wert ersetzt", "formel-fehlt", m_formel_weg),
        ("Lernbereiche-Formel verbogen", "formel-abweichung", m_lb_formel),
    ]

    print("\nMutationen (jede MUSS anschlagen):")
    schlecht = 0
    for name, regel, mut in faelle:
        wb = openpyxl.load_workbook(pfad, data_only=False)
        mut(wb)
        neu = {r for r, _ in pruefe(pfad, wb)}
        if regel in vorhanden:
            print("   ok    %-34s -> %s (schon am echten Befund)" % (name, regel))
        elif regel in neu:
            print("   ok    %-34s -> %s" % (name, regel))
        else:
            print("   FEHLT %-34s -> %s nicht gemeldet (gemeldet: %s)"
                  % (name, regel, sorted(neu) or "nichts"))
            schlecht += 1
    #  Diese eine laeuft ueber eine echte Datei statt ueber ein
    #  openpyxl-Objekt.
    ziel = m_bezug_datei(pfad)
    neu_regeln = {r for r, _ in pruefe(ziel)}
    if "bezug" in neu_regeln:
        print("   ok    %-34s -> bezug" % "Fremdblattbezug verstellt")
    else:
        print("   FEHLT %-34s -> bezug nicht gemeldet (gemeldet: %s)"
              % ("Fremdblattbezug verstellt", sorted(neu_regeln) or "nichts"))
        schlecht += 1
    os.remove(ziel)

    print("\n%d von %d Mutationen erkannt."
          % (len(faelle) + 1 - schlecht, len(faelle) + 1))
    return schlecht


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("dateien", nargs="+")
    ap.add_argument("--selbsttest", action="store_true")
    a = ap.parse_args()

    if a.selbsttest:
        return 1 if selbsttest(a.dateien[0]) else 0

    schlecht = 0
    for p in a.dateien:
        fund = [f for f in pruefe(p) if f[0] != "hinweis"]
        kurz = os.path.basename(p)
        if not fund:
            print("OK     %s" % kurz)
        else:
            schlecht += 1
            print("FEHLER %s" % kurz)
            for regel, s in fund:
                print("        [%s] %s" % (regel, s))
    return 1 if schlecht else 0


if __name__ == "__main__":
    sys.exit(main())
